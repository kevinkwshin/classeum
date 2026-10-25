"""Build the Classeum v1 quiz templates with easy English-vocabulary examples.

The main input is the 단어장 sheet: the teacher fills in a word and its meaning,
and the app generates the multiple-choice questions and wrong options. The 문항
sheet is optional, for hand-written questions the app cannot generate.

One file is one upload, and each upload becomes one dated quiz event, so the
example is split into two weekly files.

Usage: python make_vocab_template.py <output-dir>
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Malgun Gothic"
REQUIRED_FILL = PatternFill("solid", fgColor="FFE8A3")
OPTIONAL_FILL = PatternFill("solid", fgColor="E9EBEF")
TITLE_FILL = PatternFill("solid", fgColor="101218")
THIN = Side(style="thin", color="D5D8DF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LAST_ROW = 500
DEFAULT_SECONDS = 15
POS = "명,동,형,부,대,전,접,감,구"

# (header, required?, width, wrap?)
WORD_COLUMNS = [
    ("단어", True, 16, False),
    ("뜻", True, 18, False),
    ("품사", False, 7, False),
    ("헷갈리는 단어", False, 16, False),
    ("예문", False, 40, True),
    ("예문 뜻", False, 32, True),
    ("시간(초)", False, 9, False),
    ("세트", False, 12, False),
    ("단원", False, 10, False),
    ("상태", False, 8, False),
    ("단어ID", False, 10, False),
]

QUESTION_COLUMNS = [
    ("문항ID", True, 10, False),
    ("유형", True, 9, False),
    ("문제", True, 44, True),
    ("보기1", False, 14, False),
    ("보기2", False, 14, False),
    ("보기3", False, 14, False),
    ("보기4", False, 14, False),
    ("정답", True, 10, False),
    ("해설", False, 36, True),
    ("피드백1", False, 24, True),
    ("피드백2", False, 24, True),
    ("피드백3", False, 24, True),
    ("피드백4", False, 24, True),
    ("시간(초)", False, 9, False),
    ("세트", False, 12, False),
    ("상태", False, 8, False),
]

# 단어, 뜻, 품사, 헷갈리는 단어, 예문, 예문 뜻
WEEK1 = [
    ("library", "도서관", "명", "", "I study in the [library] after school.", "나는 방과 후에 도서관에서 공부한다."),
    ("borrow", "빌리다", "동", "lend", "Can I [borrow] your pen?", "네 펜 좀 빌려도 될까?"),
    ("lend", "빌려주다", "동", "borrow", "Can you [lend] me your eraser?", "지우개 좀 빌려줄래?"),
    ("subject", "과목, 주제", "명", "", "My favorite [subject] is math.", "내가 가장 좋아하는 과목은 수학이다."),
    ("quiet", "조용한", "형", "quite(꽤)", "Please be [quiet] in the library.", "도서관에서는 조용히 해 주세요."),
    ("early", "이른, 일찍", "형", "late", "I get up [early] on Mondays.", "나는 월요일에 일찍 일어난다."),
    ("late", "늦은, 늦게", "형", "early", "Don't be [late] for class.", "수업에 늦지 마."),
    ("homework", "숙제", "명", "", "I finished my [homework].", "나는 숙제를 끝냈다."),
    ("classroom", "교실", "명", "", "", ""),
    ("friend", "친구", "명", "", "", ""),
    ("test", "시험", "명", "", "We have a math [test] today.", "오늘 수학 시험이 있다."),
    ("question", "질문", "명", "answer", "Can I ask a [question]?", "질문 하나 해도 될까요?"),
    ("answer", "대답, 답", "명", "question", "I know the [answer].", "나는 답을 안다."),
    ("easy", "쉬운", "형", "difficult", "", ""),
    ("difficult", "어려운", "형", "easy", "", ""),
    ("begin", "시작하다", "동", "finish", "", ""),
    ("finish", "끝내다", "동", "begin", "", ""),
    ("practice", "연습하다", "동", "", "I [practice] English every day.", "나는 매일 영어를 연습한다."),
    ("remember", "기억하다", "동", "forget(잊다)", "", ""),
    ("break", "쉬는 시간", "명", "", "Let's play soccer at [break].", "쉬는 시간에 축구하자."),
]

WEEK2 = [
    ("weekend", "주말", "명", "weekday(주중)", "What do you do on the [weekend]?", "주말에 뭐 해?"),
    ("breakfast", "아침 식사", "명", "dinner", "I eat [breakfast] at seven.", "나는 7시에 아침을 먹는다."),
    ("dinner", "저녁 식사", "명", "breakfast", "", ""),
    ("clothes", "옷", "명", "cloth(천)", "I bought new [clothes].", "나는 새 옷을 샀다."),
    ("tired", "피곤한", "형", "", "I am [tired] today.", "나는 오늘 피곤하다."),
    ("hungry", "배고픈", "형", "angry(화난)", "", ""),
    ("listen", "듣다", "동", "hear(들리다)", "[Listen] to the teacher.", "선생님 말씀을 잘 들어라."),
    ("their", "그들의", "대", "there(거기에)", "This is [their] house.", "이것은 그들의 집이다."),
    ("teeth", "이 (여러 개)", "명", "tooth(이 하나)", "I brush my [teeth] every morning.", "나는 매일 아침 이를 닦는다."),
    ("kitchen", "부엌", "명", "chicken(닭)", "", ""),
    ("wash", "씻다", "동", "", "", ""),
    ("sleep", "자다", "동", "", "", ""),
    ("wake", "깨다", "동", "", "", ""),
    ("cook", "요리하다", "동", "", "My dad [cooks] dinner.", "아빠가 저녁을 요리하신다."),
    ("shower", "샤워", "명", "", "", ""),
    ("morning", "아침", "명", "evening", "", ""),
    ("evening", "저녁", "명", "morning", "", ""),
    ("always", "항상", "부", "never(절대 ~않다)", "", ""),
    ("sometimes", "가끔", "부", "always", "", ""),
    ("busy", "바쁜", "형", "", "", ""),
]

# 문항ID, 유형, 문제, 보기1-4, 정답, 해설, 피드백1-4
WEEK1_QUESTIONS = [
    ("Q001", "OX", "homework는 셀 수 있는 명사라서 homeworks라고 쓸 수 있다.", "", "", "", "", "X",
     "homework는 셀 수 없는 명사라서 -s를 붙이지 않아요.", "", "", "", ""),
    ("Q002", "객관식", "빈칸에 알맞은 것은? My sister ___ me her bike yesterday.", "lent", "borrowed", "bought", "brought", "1",
     "lend(빌려주다)의 과거형은 lent예요.", "", "borrow는 '빌리다'라서 주는 사람에게 쓰지 않아요.", "bought는 '샀다'예요.", "brought는 '가져왔다'예요."),
]

GUIDE = [
    ("Classeum 문항 템플릿 v1", None),
    ("", None),
    ("가장 쉬운 방법: 단어장 시트", "h"),
    ("'단어장' 시트에 단어와 뜻 두 칸만 채우세요. 객관식 문제와 오답 보기는 앱이 만들어요.", None),
    ("앱이 만드는 문제: 뜻 고르기(borrow → 빌리다), 단어 고르기(빌리다 → borrow), 예문 빈칸(예문이 있을 때)", None),
    ("오답 보기는 같은 단어장의 다른 단어에서 고르고, 헷갈리는 단어 칸에 쓴 단어를 가장 먼저 씁니다.", None),
    ("틀린 학생에게는 '네가 고른 lend는 빌려주다예요'처럼 피드백이 자동으로 나가요.", None),
    ("", None),
    ("칸 설명", "h"),
    ("단어·뜻 (필수): 뜻이 여러 개면 쉼표로 나눕니다. (예: 과목, 주제)", None),
    ("품사 (선택): 명·동·형·부·대·전·접·감·구 중에서 고르면 비슷한 오답을 고르는 데 써요.", None),
    ("헷갈리는 단어 (선택): 여러 개면 ; 로 나눕니다. 단어장에 없는 단어는 괄호에 뜻을 씁니다. (예: quite(꽤))", None),
    ("예문 (선택): 빈칸이 될 자리를 [ ] 로 감쌉니다. (예: Can I [borrow] your pen?)", None),
    (f"시간(초) (선택): 비워 두면 {DEFAULT_SECONDS}초입니다. 특별히 오래 걸리는 단어만 숫자를 쓰세요.", None),
    ("세트 (선택): 비워 두면 올린 날짜로 이벤트가 생겨요. (예: 10-29 퀴즈)", None),
    ("상태 (선택): 사용(기본) · 검토(저장만 하고 출제 안 함) · 보관(숨김)", None),
    ("단어ID (선택): 거의 쓰지 않아요. 철자는 같고 뜻이 다른 단어를 따로 기록할 때만 씁니다.", None),
    ("", None),
    ("기록이 이어지는 방식", "h"),
    ("같은 단어를 다음 주에 다시 올리면 같은 단어로 이어져서, 학생별로 자주 틀리는 단어가 쌓여요.", None),
    ("", None),
    ("문항 시트 (선택)", "h"),
    ("반의어, 문법처럼 단어장으로 만들 수 없는 문제만 직접 씁니다. 유형은 객관식 / OX / 단답형.", None),
    ("객관식 정답은 번호(1~4), OX는 O 또는 X, 단답형은 답 (여러 답은 a|b).", None),
    ("", None),
    ("명단", "h"),
    ("학생 명단은 이 파일에 넣지 않아요. 선생님 화면 [학생·기기]에서 반을 만들고, 반마다 '번호 이름'을 한 번만 붙여 넣으세요.", None),
]



def style_cell(cell, *, bold=False, fill=None, wrap=False, color="101218", size=11, center=False):
    cell.font = Font(name=FONT, bold=bold, color=color, size=size)
    if fill:
        cell.fill = fill
    cell.alignment = Alignment(vertical="center", wrap_text=wrap, horizontal="center" if center else None)


def write_header(ws, columns):
    for col, (header, required, width, _) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=col, value=header)
        style_cell(c, bold=True, fill=REQUIRED_FILL if required else OPTIONAL_FILL, center=True)
        c.border = BOX
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "C2" if columns is WORD_COLUMNS else "B2"


def write_rows(ws, columns, rows):
    for r, row in enumerate(rows, start=2):
        for col, value in enumerate(row, start=1):
            c = ws.cell(row=r, column=col, value=value if value != "" else None)
            style_cell(c, wrap=columns[col - 1][3])
            c.border = BOX


def add_list(ws, columns, header, options, error):
    col = [h for h, *_ in columns].index(header) + 1
    letter = ws.cell(row=1, column=col).column_letter
    dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True,
                        showErrorMessage=True, errorTitle=header, error=error)
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}{LAST_ROW}")


def add_seconds_rule(ws, columns):
    col = [h for h, *_ in columns].index("시간(초)") + 1
    letter = ws.cell(row=1, column=col).column_letter
    dv = DataValidation(type="whole", operator="between", formula1="5", formula2="60", allow_blank=True,
                        showErrorMessage=True, errorTitle="시간(초)", error="5~60 사이의 숫자를 쓰거나 비워 두세요.")
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}{LAST_ROW}")


def set_text_format(ws, columns, headers):
    names = [h for h, *_ in columns]
    for header in headers:
        col = names.index(header) + 1
        for r in range(2, LAST_ROW + 1):
            ws.cell(row=r, column=col).number_format = "@"


def build(path, words, questions):
    wb = Workbook()

    guide = wb.active
    guide.title = "안내"
    guide.column_dimensions["A"].width = 110
    for i, (text, kind) in enumerate(GUIDE, start=1):
        c = guide.cell(row=i, column=1, value=text)
        if i == 1:
            style_cell(c, bold=True, fill=TITLE_FILL, color="FFFFFF", size=14)
            guide.row_dimensions[i].height = 28
        elif kind == "h":
            style_cell(c, bold=True, size=12)
            guide.row_dimensions[i].height = 22
        else:
            style_cell(c)
    guide.sheet_view.showGridLines = False

    ws = wb.create_sheet("단어장")
    write_header(ws, WORD_COLUMNS)
    write_rows(ws, WORD_COLUMNS, [w + ("", "", "", "", "") for w in words])
    add_list(ws, WORD_COLUMNS, "품사", POS, "명, 동, 형, 부, 대, 전, 접, 감, 구 중에서 고르세요.")
    add_list(ws, WORD_COLUMNS, "상태", "사용,검토,보관", "사용, 검토, 보관 중에서 고르세요.")
    add_seconds_rule(ws, WORD_COLUMNS)
    ws.auto_filter.ref = f"A1:K{len(words) + 1}"

    qs = wb.create_sheet("문항")
    write_header(qs, QUESTION_COLUMNS)
    write_rows(qs, QUESTION_COLUMNS, [q + ("", "", "") for q in questions])
    add_list(qs, QUESTION_COLUMNS, "유형", "객관식,OX,단답형", "객관식, OX, 단답형 중에서 고르세요.")
    add_list(qs, QUESTION_COLUMNS, "상태", "사용,검토,보관", "사용, 검토, 보관 중에서 고르세요.")
    add_seconds_rule(qs, QUESTION_COLUMNS)
    # Text format: 정답 keeps "1" and "O" as typed, and IDs, labels and short options like "3-1" or "7:30"
    # never become dates or times. (시간(초) stays numeric for its 5-60 rule.)
    set_text_format(ws, WORD_COLUMNS, ["세트", "단원", "단어ID"])
    set_text_format(qs, QUESTION_COLUMNS, ["문항ID", "보기1", "보기2", "보기3", "보기4", "정답", "세트"])


    wb.active = 1
    wb.save(path)


if __name__ == "__main__":
    out = sys.argv[1]
    os.makedirs(out, exist_ok=True)
    build(os.path.join(out, "classeum-vocab-week1.xlsx"), WEEK1, WEEK1_QUESTIONS)
    build(os.path.join(out, "classeum-vocab-week2.xlsx"), WEEK2, [])
